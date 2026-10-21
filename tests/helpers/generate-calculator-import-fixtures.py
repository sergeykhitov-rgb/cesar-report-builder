#!/usr/bin/env python3
"""Create minimal deterministic calculator-import workbooks for integration tests."""

from pathlib import Path
import sys

from openpyxl import Workbook


def save_reference(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    tables = {
        "README": [
            ["Файл", "Fixture"],
            ["Назначение", "Integration test"],
        ],
        "B2C_Тарифы": [],
        "B2C_Тарифы1": [
            ["Категория", "Регион", "Тип объекта", "Тариф", "Код тарифа", "АП, ₽/мес.", "Видео", "Бытовые датчики", "Расширенная ответственность", "Состав тарифа", "Источник"],
            ["B2C", "Тест", "Квартира", "Базовый", "TEST", 1200, "Нет", "Нет", "Нет", "Мониторинг", "Fixture"],
        ],
        "B2C_Переходы": [
            ["Сегмент/объект", "Исходный тариф", "Условие", "Результат", "Автоматический переход", "Примечание"],
        ],
        "B2C_Аренда": [
            ["Категория", "Регион", "Тип объекта", "Тариф", "Код тарифа", "АП аренда, ₽/мес."],
            ["B2C", "Тест", "Квартира", "Базовый", "TEST-R", 1400],
        ],
        "B2B_Тарифы": [
            ["Категория", "Регион", "Уровень цены", "Название тарифа", "Состав", "Стоимость АП, ₽/мес.", "Аренда"],
            ["B2B", "Тест", "Рекомендуемый", "Бизнес", "ОС", 3000, "Нет"],
        ],
        "Опции": [
            ["Категория", "Опция", "Периодичность", "Стоимость, ₽/мес.", "Условие/примечание", "Источник"],
            ["B2C", "Fixture option", "Ежемесячная", 250, "", "Fixture"],
        ],
        "Аренда_Логика": [["Категория", "Сценарий", "Параметр", "Значение", "Примечание"]],
        "Мат_ответственность": [
            ["Категория", "Тип объекта", "комплект абонентская плата", "материальная ответственность"],
            ["B2C", "Квартира", "базовый", 1000000],
        ],
        "Требует_решения": [["Блок", "Тема", "Что известно", "Статус"]],
        "Скидки": [["Скидка, %", "Статус/правило"], [0, "Fixture"]],
        "Время_реагирования": [["Тип территории", "Норматив", "Ручная корректировка"]],
    }
    for title, rows in tables.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    workbook.save(path)


def save_equipment(path: Path, quantity: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ритм fixture"
    # Deliberately reordered, using accepted source wording variants.
    headers = [
        "Модель",
        "Цена оборудования с установкой за единицу, руб. с НДС",
        "Ед. изм.",
        "Наименование оборудования",
        "Кол-во",
        "Код",
    ]
    sheet.append(headers)
    sheet.append(["FX-1", 1234, "шт.", "Fixture sensor", quantity, "FX-001"])
    # Real specification sheets may repeat their table header between sections.
    sheet.append(headers)
    # Signature/footer rows may contain text under quantity but have no item code.
    sheet.append([None, None, None, None, "ИСПОЛНИТЕЛЬ", None])
    workbook.save(path)


def main() -> None:
    output = Path(sys.argv[1])
    output.mkdir(parents=True, exist_ok=True)
    save_reference(output / "reference.xlsx")
    save_equipment(output / "template.xlsx", 0)
    save_equipment(output / "example.xlsx", 2)


if __name__ == "__main__":
    main()
