#!/usr/bin/env python3
"""Import CESAR calculator workbooks into stable, reviewable JSON catalogs."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from typing import Any, Iterable

import xlrd
from openpyxl import load_workbook


SYSTEM_LABELS = {
    "ritm": "Ритм",
    "si-nord": "Си-Норд",
    "hikvision": "Hikvision",
    "rental-used": "Б/У комплекты для аренды",
    "other": "Другое",
}
STOCK_SHEETS = {"13.08.26", "13.08.26 2725"}

# These inventory codes describe the same sellable item in several specification
# sheets, but the workbook editors used annotations or omitted a model in some
# copies. Only the exact reviewed variants below may share a stable key. A new
# spelling or model is treated as a conflict instead of being silently merged.
EQUIPMENT_VARIANT_ALLOWLIST = {
    ("si-nord", "91068"): {
        "canonical": ("Комплект Выгодный дом Си-Норд (панель, клав, СМК, 2ИК, Брелок, 1ДУВ, ТЕМП, ДИП, АНТ)", "Норд-GSM"),
        "variants": {
            ("Комплект Выгодный дом Си-Норд (панель, клав, СМК, 2ИК, Брелок, ДУВ, ТЕМП, ДИП, АНТ)", "Норд-GSM"),
            ("Комплект Выгодный дом Си-Норд (панель, клав, СМК, 2ИК, Брелок, 1ДУВ, ТЕМП, ДИП, АНТ)", "Норд-GSM"),
        },
    },
    ("si-nord", "64495"): {
        "canonical": ("Адаптер СИ-НОРД, Ethemet", "Ethemet"),
        "variants": {(я, м) for я, м in [
            ("Адаптер СИ-НОРД, Ethemet", ""),
            ("Адаптер СИ-НОРД, Ethemet", "Ethemet"),
        ]},
    },
    ("si-nord", "98957"): {
        "canonical": ("Ящик с кодовым механизмом - сейф-бокс для ключей", "Сейф-бокс"),
        "variants": {(я, м) for я, м in [
            ("Ящик с кодовым механизмом - сейф-бокс для ключей", ""),
            ("Ящик с кодовым механизмом - сейф-бокс для ключей", "Сейф-бокс"),
        ]},
    },
    ("hikvision", "108840"): {
        "canonical": ("Датчик ИК беспроводной, защита от животных Hikvision, DS-PDP15P-EG2-WE,", "PIR"),
        "variants": {(я, "PIR") for я in [
            "Датчик ИК беспроводной, защита от животных Hikvision, DS-PDP15P-EG2-WE,",
            "Датчик ИК беспроводной, защита от животных Hikvision, DS-PDP15P-EG2-WE, ТРИ ВАРИАНТА ЗАМЕНЫ НИЖЕ",
        ]},
    },
    ("hikvision", "103206"): {
        "canonical": ("Датчик протечки воды (AX PRO) Hikvision, DS-PDWL-E-WE,", "WaterLeak"),
        "variants": {(я, "WaterLeak") for я in [
            "Датчик протечки воды (AX PRO) Hikvision, DS-PDWL-E-WE,",
            "Датчик протечки воды (AX PRO) Hikvision, DS-PDWL-E-WE, ВАРИАНТ ЗАМЕНЫ НИЖЕ",
        ]},
    },
    ("hikvision", "108841"): {
        "canonical": ("Извещатель пожар, дымовой Hikvision, DS-PDSMK-S-WE,", "Smoke"),
        "variants": {(я, "Smoke") for я in [
            "Извещатель пожар, дымовой Hikvision, DS-PDSMK-S-WE,",
            "Извещатель пожар, дымовой Hikvision, DS-PDSMK-S-WE, вариант для панелей RU ниже",
        ]},
    },
    ("hikvision", "108844"): {
        "canonical": ("Датчик СМК беспроводной Hikvision, DS-PDMC-EG2-WE,", "Door"),
        "variants": {(я, "Door") for я in [
            "Датчик СМК беспроводной Hikvision, DS-PDMC-EG2-WE,",
            "Датчик СМК беспроводной Hikvision, DS-PDMC-EG2-WE, ВАРИАНТ ЗАМЕНЫ НИЖЕ",
        ]},
    },
    ("si-nord", "89012"): {
        "canonical": ("Датчик температурный СИ-НОРД, выносной, проводной", "выносной"),
        "variants": {(я, м) for я, м in [
            ("Датчик температурный СИ-НОРД, выносной, проводной", ""),
            ("Датчик температурный СИ-НОРД, выносной, проводной", "выносной"),
        ]},
    },
    ("si-nord", "100334"): {
        "canonical": ("Расширитель СИ-НОРД охранный, ЕХ81", "ЕХ81"),
        "variants": {(
            "Расширитель СИ-НОРД, ЕХ81 проводных шлейфов", ""
        ), (
            "Расширитель СИ-НОРД охранный, ЕХ81", "ЕХ81"
        )},
    },
    ("si-nord", "102022"): {
        "canonical": ("Расширитель СИ-НОРД пожарный, ЕХ22", "ЕХ22"),
        "variants": {(
            "Расширитель СИ-НОРД, ЕХ22 пожарных шлейфов", ""
        ), (
            "Расширитель СИ-НОРД пожарный, ЕХ22", "ЕХ22"
        )},
    },
    ("si-nord", "64491"): {
        "canonical": ("Датчик магнитоконтактный СИ-НОРД, СН-СМК исп.1", "СН-СМК"),
        "variants": {(я, "СН-СМК") for я in [
            "Датчик магнитоконтактный СИ-НОРД, СН-СМК исп.1",
            "Датчик магнитоконтактный СИ-НОРД, СН-СМК исп.1 (большой,ПРД для подключения стороннего оборудования)",
        ]},
    },
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_number(value: Any, *, allow_none: bool = False) -> int | float | None:
    if value is None or normalize_text(value) == "":
        if allow_none:
            return None
        raise ValueError("numeric value is blank")
    if isinstance(value, str):
        value = value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
        value = re.sub(r"[^0-9.\-]", "", value)
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        if allow_none:
            return None
        raise ValueError(f"not a numeric value: {value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"non-finite numeric value: {value!r}")
    return int(number) if number.is_integer() else number


def normalize_code(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = normalize_text(value)
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def source_ref(workbook: Path, sheet: str, row: int) -> dict[str, Any]:
    return {"workbook": workbook.name, "sheet": sheet, "row": row}


def source_with_fields(
    workbook: Path,
    sheet: str,
    row: int,
    **fields: Any,
) -> dict[str, Any]:
    return {**source_ref(workbook, sheet, row), **fields}


def table_rows(ws: Any) -> list[tuple[int, list[Any]]]:
    return [(index, list(row)) for index, row in enumerate(ws.iter_rows(values_only=True), 1)]


def records_from_rows(
    workbook: Path,
    sheet: str,
    rows: list[tuple[int, list[Any]]],
    field_names: list[str],
) -> list[dict[str, Any]]:
    records = []
    for row_number, values in rows[1:]:
        padded = values + [None] * (len(field_names) - len(values))
        if not any(normalize_text(value) for value in padded[: len(field_names)]):
            continue
        record = {
            field: normalize_text(value)
            for field, value in zip(field_names, padded)
        }
        record["source"] = source_ref(workbook, sheet, row_number)
        records.append(record)
    return records


def read_calculator_reference(reference_path: Path) -> dict[str, Any]:
    workbook = load_workbook(reference_path, data_only=True, read_only=True)
    sheets = {ws.title: table_rows(ws) for ws in workbook.worksheets}

    def rows(name: str) -> list[tuple[int, list[Any]]]:
        if name not in sheets:
            raise ValueError(f"required reference sheet is missing: {name}")
        return sheets[name]

    b2c_tariffs = []
    for row_number, values in rows("B2C_Тарифы1")[1:]:
        if not normalize_text(values[4] if len(values) > 4 else None):
            continue
        b2c_tariffs.append({
            "category": normalize_text(values[0]),
            "region": normalize_text(values[1]),
            "objectType": normalize_text(values[2]),
            "name": normalize_text(values[3]),
            "code": normalize_code(values[4]),
            "monthlyPrice": normalize_number(values[5]),
            "hasVideo": normalize_text(values[6]).lower() == "да",
            "hasHouseholdSensors": normalize_text(values[7]).lower() == "да",
            "hasExtendedLiability": normalize_text(values[8]).lower() == "да",
            "contents": normalize_text(values[9]),
            "authority": normalize_text(values[10]),
            "source": source_ref(reference_path, "B2C_Тарифы1", row_number),
        })

    b2c_rental = []
    for row_number, values in rows("B2C_Аренда")[1:]:
        if not normalize_text(values[4] if len(values) > 4 else None):
            continue
        b2c_rental.append({
            "category": normalize_text(values[0]),
            "region": normalize_text(values[1]),
            "objectType": normalize_text(values[2]),
            "name": normalize_text(values[3]),
            "code": normalize_code(values[4]),
            "monthlyPrice": normalize_number(values[5]),
            "source": source_ref(reference_path, "B2C_Аренда", row_number),
        })

    b2b_tariffs = []
    for row_number, values in rows("B2B_Тарифы")[1:]:
        if not normalize_text(values[3] if len(values) > 3 else None):
            continue
        b2b_tariffs.append({
            "category": normalize_text(values[0]),
            "region": normalize_text(values[1]),
            "priceLevel": normalize_text(values[2]),
            "name": normalize_text(values[3]),
            "contents": normalize_text(values[4]),
            "monthlyPrice": normalize_number(values[5]),
            "rental": normalize_text(values[6]).lower() == "да",
            "source": source_ref(reference_path, "B2B_Тарифы", row_number),
        })

    options = []
    for row_number, values in rows("Опции")[1:]:
        name = normalize_text(values[1] if len(values) > 1 else None)
        if not name:
            continue
        price = normalize_number(values[3], allow_none=True)
        options.append({
            "category": normalize_text(values[0]),
            "name": name,
            "frequency": normalize_text(values[2]),
            "monthlyPrice": price,
            "pricing": "fixed" if price is not None else "individual",
            "condition": normalize_text(values[4]),
            "authority": normalize_text(values[5]),
            "source": source_ref(reference_path, "Опции", row_number),
        })

    liability_entries = []
    liability = {"b2c": {"apartment": {}, "cottage": {}}, "b2b": {"any": {}}}
    level_ids = {"базовый": "basic", "выгодный": "value", "премиальный": "premium", "любой": "any"}
    object_ids = {"квартира": "apartment", "коттедж": "cottage", "любой": "any"}
    for row_number, values in rows("Мат_ответственность")[1:]:
        category = normalize_text(values[0]).lower()
        object_label = normalize_text(values[1])
        level_label = normalize_text(values[2])
        amount = normalize_number(values[3])
        object_id = object_ids[object_label.lower()]
        level_id = level_ids[level_label.lower()]
        segment_id = category.lower()
        liability[segment_id][object_id][level_id] = amount
        liability_entries.append({
            "category": normalize_text(values[0]),
            "objectType": object_label,
            "level": level_label,
            "amount": amount,
            "source": source_ref(reference_path, "Мат_ответственность", row_number),
        })
    liability["entries"] = liability_entries

    discounts = []
    for row_number, values in rows("Скидки")[1:]:
        label = normalize_text(values[0] if values else None)
        if not label:
            continue
        value = normalize_number(values[0], allow_none=True)
        discounts.append({
            "percent": value,
            "kind": "standard" if value is not None else "individual",
            "label": label,
            "rule": normalize_text(values[1]),
            "source": source_ref(reference_path, "Скидки", row_number),
        })

    readme = {
        normalize_text(values[0]): normalize_text(values[1])
        for _, values in rows("README")
        if len(values) >= 2 and normalize_text(values[0])
    }

    return {
        "schemaVersion": 1,
        "sourceWorkbook": reference_path.name,
        "readme": readme,
        "b2cTariffs": b2c_tariffs,
        "b2cTransitions": records_from_rows(reference_path, "B2C_Переходы", rows("B2C_Переходы"), [
            "segmentObject", "sourceTariff", "condition", "result", "automaticTransition", "note"
        ]),
        "b2cRentalTariffs": b2c_rental,
        "b2bTariffs": b2b_tariffs,
        "options": options,
        "rentalLogic": records_from_rows(reference_path, "Аренда_Логика", rows("Аренда_Логика"), [
            "category", "scenario", "parameter", "value", "note"
        ]),
        "materialLiability": liability,
        "discounts": discounts,
        "responseTimes": records_from_rows(reference_path, "Время_реагирования", rows("Время_реагирования"), [
            "territoryType", "target", "manualAdjustment"
        ]),
        "unresolvedDecisions": records_from_rows(reference_path, "Требует_решения", rows("Требует_решения"), [
            "block", "topic", "known", "status"
        ]),
    }


def system_for_sheet(sheet: str) -> str:
    folded = sheet.casefold()
    if "б_у" in folded or "б/у" in folded or "аренд" in folded:
        return "rental-used"
    if "hikvision" in folded:
        return "hikvision"
    if "ритм" in folded:
        return "ritm"
    if "си-норд" in folded or "си_норд" in folded or "неотапливаемые" in folded:
        return "si-nord"
    return "other"


def find_specification_header(sheet: Any) -> int | None:
    for row_index in range(sheet.nrows):
        values = [normalize_text(value).casefold() for value in sheet.row_values(row_index)]
        if "код" in values and any("наименование оборудования" in value for value in values):
            return row_index
    return None


def read_equipment_rows(workbook_path: Path) -> tuple[list[dict[str, Any]], dict[str, int]]:
    workbook = xlrd.open_workbook(workbook_path, on_demand=True)
    rows: list[dict[str, Any]] = []
    sheet_counts: dict[str, int] = {}
    for sheet_name in workbook.sheet_names():
        if sheet_name in STOCK_SHEETS:
            continue
        sheet = workbook.sheet_by_name(sheet_name)
        header_index = find_specification_header(sheet)
        if header_index is None:
            continue
        count = 0
        for row_index in range(header_index + 1, sheet.nrows):
            values = sheet.row_values(row_index)
            code = normalize_code(values[1] if len(values) > 1 else None)
            name = normalize_text(values[2] if len(values) > 2 else None)
            if not code or not name:
                continue
            price = normalize_number(values[6] if len(values) > 6 else None, allow_none=True)
            if price is None:
                continue
            rows.append({
                "system": system_for_sheet(sheet_name),
                "systemLabel": SYSTEM_LABELS[system_for_sheet(sheet_name)],
                "code": code,
                "name": name,
                "model": normalize_text(values[3] if len(values) > 3 else None),
                "unit": normalize_text(values[4] if len(values) > 4 else None),
                "installedUnitPrice": price,
                "source": source_with_fields(
                    workbook_path,
                    sheet_name,
                    row_index + 1,
                    name=name,
                    model=normalize_text(values[3] if len(values) > 3 else None),
                ),
            })
            count += 1
        sheet_counts[sheet_name] = count
    workbook.release_resources()
    return rows, sheet_counts


def deduplicate_equipment(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    items: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        key = (row["system"], row["code"])
        allowed = EQUIPMENT_VARIANT_ALLOWLIST.get(key)
        variant = (row["name"], row["model"])
        current = items.get(key)
        if current is None:
            if allowed and variant not in allowed["variants"]:
                source = row["source"]
                raise ValueError(
                    f"unreviewed equipment variant {key!r} in "
                    f"{source['workbook']} / {source['sheet']} row {source['row']}: name, model"
                )
            item = {key: value for key, value in row.items() if key != "source"}
            if allowed:
                item["name"], item["model"] = allowed["canonical"]
            item["sources"] = [row["source"]]
            items[key] = item
            continue
        commercial_fields = ("name", "model", "unit", "installedUnitPrice")
        conflicts = [field for field in commercial_fields if current[field] != row[field]]
        if allowed and variant in allowed["variants"]:
            conflicts = [field for field in conflicts if field not in ("name", "model")]
        if conflicts:
            first_source = current["sources"][0]
            source = row["source"]
            raise ValueError(
                f"conflicting duplicate equipment key {key!r}: {', '.join(conflicts)}; "
                f"first source {first_source['workbook']} / {first_source['sheet']} "
                f"row {first_source['row']}; conflicting source "
                f"{source['workbook']} / {source['sheet']} row {source['row']}"
            )
        provenance = row["source"]
        if provenance not in current["sources"]:
            current["sources"].append(provenance)
        if not current["model"] and row["model"]:
            current["model"] = row["model"]
    return sorted(items.values(), key=lambda item: (item["system"], item["code"]))


def read_example_control(example_path: Path) -> dict[str, Any]:
    rows, sheet_counts = read_equipment_rows(example_path)
    workbook = xlrd.open_workbook(example_path, on_demand=True)
    selected = []
    for sheet_name in workbook.sheet_names():
        if sheet_name in STOCK_SHEETS:
            continue
        sheet = workbook.sheet_by_name(sheet_name)
        header_index = find_specification_header(sheet)
        if header_index is None:
            continue
        for row_index in range(header_index + 1, sheet.nrows):
            values = sheet.row_values(row_index)
            code = normalize_code(values[1] if len(values) > 1 else None)
            quantity = normalize_number(values[5] if len(values) > 5 else None, allow_none=True)
            if code and quantity is not None and quantity > 0:
                selected.append({
                    "sheet": sheet_name,
                    "row": row_index + 1,
                    "code": code,
                    "quantity": quantity,
                })
    workbook.release_resources()
    return {
        "workbook": example_path.name,
        "role": "control-and-provenance-only",
        "specificationRowCount": len(rows),
        "rowsBySheet": sheet_counts,
        "selectedRows": selected,
    }


def read_equipment_template(template_path: Path, example_path: Path) -> dict[str, Any]:
    rows, sheet_counts = read_equipment_rows(template_path)
    return {
        "schemaVersion": 1,
        "sourceWorkbook": template_path.name,
        "excludedSheets": sorted(STOCK_SHEETS.intersection(workbook_sheet_names(template_path))),
        "rowsBySheet": sheet_counts,
        "items": deduplicate_equipment(rows),
        "exampleControl": read_example_control(example_path),
    }


def validate_calculator(calculator: dict[str, Any]) -> None:
    required = (
        "b2cTariffs", "b2cTransitions", "b2cRentalTariffs", "b2bTariffs", "options",
        "rentalLogic", "materialLiability", "discounts", "responseTimes", "unresolvedDecisions",
    )
    missing = [key for key in required if key not in calculator]
    if missing:
        raise ValueError(f"calculator catalog is missing sections: {', '.join(missing)}")
    priced_rows = calculator["b2cTariffs"] + calculator["b2cRentalTariffs"] + calculator["b2bTariffs"]
    if not priced_rows or any(not isinstance(row["monthlyPrice"], (int, float)) for row in priced_rows):
        raise ValueError("calculator tariff prices must be numeric")


def validate_equipment(equipment: dict[str, Any]) -> None:
    seen: set[tuple[str, str]] = set()
    for item in equipment["items"]:
        key = (item["system"], item["code"])
        if key in seen:
            raise ValueError(f"duplicate equipment key after normalization: {key!r}")
        seen.add(key)
        if not item["name"] or not isinstance(item["installedUnitPrice"], (int, float)):
            raise ValueError(f"invalid equipment item: {key!r}")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_sheet_names(path: Path) -> list[str]:
    if path.suffix.casefold() == ".xlsx":
        workbook = load_workbook(path, read_only=True)
        names = list(workbook.sheetnames)
        workbook.close()
        return names
    workbook = xlrd.open_workbook(path, on_demand=True)
    names = workbook.sheet_names()
    workbook.release_resources()
    return names


def build_manifest(paths: Iterable[Path]) -> dict[str, Any]:
    return {
        "schemaVersion": 1,
        "deterministic": True,
        "sources": [
            {
                "basename": path.name,
                "sha256": sha256(path),
                "sizeBytes": path.stat().st_size,
                "sheetNames": workbook_sheet_names(path),
            }
            for path in paths
        ],
    }


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def import_sources(
    reference_path: Path,
    template_path: Path,
    filled_example_path: Path,
    output_dir: Path,
) -> tuple[dict[str, Any], dict[str, Any]]:
    calculator = read_calculator_reference(reference_path)
    equipment = read_equipment_template(template_path, filled_example_path)
    validate_calculator(calculator)
    validate_equipment(equipment)
    write_json(output_dir / "calculator.json", calculator)
    write_json(output_dir / "equipment.json", equipment)
    write_json(output_dir / "import-manifest.json", build_manifest([
        reference_path, template_path, filled_example_path
    ]))
    return calculator, equipment


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reference", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--example", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    calculator, equipment = import_sources(args.reference, args.template, args.example, args.output)
    summary = {
        "calculatorRows": {
            key: len(calculator[key])
            for key in (
                "b2cTariffs", "b2cTransitions", "b2cRentalTariffs", "b2bTariffs", "options",
                "rentalLogic", "discounts", "responseTimes", "unresolvedDecisions",
            )
        },
        "equipmentRowsBySheet": equipment["rowsBySheet"],
        "equipmentItems": len(equipment["items"]),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
